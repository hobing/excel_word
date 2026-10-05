# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches,Cm
import os

#  'C:\\Users\\oobooH\\Desktop\\005\\开发区治安监控竣工信息表重序02.xlsx'


'''
import warnings
warnings.filterwarnings('ignore')
'''

doc01 = Document('C:\\Users\\ooborh\\Desktop\\001\\001.docx')
#doc02 = Document()
'''
for para in doc01.paragraphs:
        found = para.text.find('123 ')
        if found != -1:
            #para.text = para.text.replace('123', '456')
            n_pra = para.insert_paragraph_before(' ')
            
# 将段落内容清空
            paragraph.clear()
            
# 存储需要删除的段落索引
    paragraphs_to_delete = []
    
    # 遍历所有段落，找出包含指定文本的段落
    for i, paragraph in enumerate(doc.paragraphs):
        if text_to_delete in paragraph.text:
            paragraphs_to_delete.append(i)
    
    # 从后往前删除段落，避免索引变化问题
    for i in reversed(paragraphs_to_delete):
        # 获取段落元素并删除
        p = doc.paragraphs[i]._p
        p.getparent().remove(p)
    
    # 保存修改后的文档
    doc.save(file_path)

for text in  doc01.paragraphs[0].text.split(' '):
    doc02.add_paragraph(text)
'''
'''
#遍历文档所有单元格，单元格是高于段落的，一个单元格可以有多个段落
for table in doc01.tables:
    for row in table.rows:
        for cell in row.cells:
            if '_A1_' in cell.text:
                run = cell.paragraphs[0].add_run()

                # 单位是英寸Inches，注意import内容
                run.add_picture('C:\\Users\\ooborh\\Desktop\\001\\184.PNG',
                                width=Inches(3), height=Cm(7.6))
                new_p = cell.add_paragraph()
                run.add_picture('C:\\Users\\ooborh\\Desktop\\001\\184.PNG',
                                width=Inches(3), height=Cm(7.6))
'''


'''
for para in doc01.paragraphs:
        found = para.text.find('_A1_')
        if found != -1:
            run = para.add_run()

            #单位是英寸Inches，注意import内容
            run.add_picture('C:\\Users\\ooborh\\Desktop\\001\\184.PNG',
                            width =Inches(3), height =Cm(7.6))
'''
doc01.save('C:\\Users\\ooborh\\Desktop\\001\\002.docx')

